from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "platform" / "docs" / "OPERATOR_END_TO_END_RUNBOOK.md"
OUTPUT = ROOT.parents[1] / "Design_Documents_DB_First" / "CbCR_End_to_End_Operator_Runbook.docx"

PAGE_WIDTH = 12240
PAGE_HEIGHT = 15840
MARGIN = 1440
USABLE_DXA = 9360
TABLE_INDENT = 120
CELL_MARGIN_TOP = 80
CELL_MARGIN_BOTTOM = 80
CELL_MARGIN_START = 120
CELL_MARGIN_END = 120

NAVY = "17365D"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
MUTED = "64748B"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
CALLOUT = "F4F6F9"
BORDER = "B8C4D1"
WHITE = "FFFFFF"
BLACK = "1F2933"


def set_run_font(run, name="Calibri", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (
        ("top", CELL_MARGIN_TOP), ("bottom", CELL_MARGIN_BOTTOM),
        ("start", CELL_MARGIN_START), ("end", CELL_MARGIN_END),
    ):
        tag = tc_mar.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            tc_mar.append(tag)
        tag.set(qn("w:w"), str(value))
        tag.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    if sum(widths) != USABLE_DXA:
        raise ValueError(f"table widths must total {USABLE_DXA}: {widths}")
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(USABLE_DXA))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT))
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = Inches(widths[index] / 1440)
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tc_w)
            tc_w.set(qn("w:w"), str(widths[index]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table_borders(table, color=BORDER, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    flag = OxmlElement("w:tblHeader")
    flag.set(qn("w:val"), "true")
    tr_pr.append(flag)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def set_keep_with_next(paragraph, value=True):
    p_pr = paragraph._p.get_or_add_pPr()
    node = p_pr.find(qn("w:keepNext"))
    if value and node is None:
        p_pr.append(OxmlElement("w:keepNext"))
    elif not value and node is not None:
        p_pr.remove(node)


def shade_paragraph(paragraph, fill=CALLOUT, border=BORDER):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    for edge in ("top", "left", "bottom", "right"):
        line = OxmlElement(f"w:{edge}")
        line.set(qn("w:val"), "single")
        line.set(qn("w:sz"), "4")
        line.set(qn("w:space"), "4")
        line.set(qn("w:color"), border)
        p_bdr.append(line)


def add_field(paragraph, instruction, fallback="1"):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = fallback
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instr, separate, text, end):
        run._r.append(node)


def add_numbering(doc, marker_type):
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "decimal" if marker_type == "decimal" else "bullet")
    level.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText")
    marker = "%1." if marker_type == "decimal" else ("□" if marker_type == "check" else "•")
    lvl_text.set(qn("w:val"), marker)
    level.append(lvl_text)
    suff = OxmlElement("w:suff")
    suff.set(qn("w:val"), "tab")
    level.append(suff)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    p_pr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "271")
    p_pr.append(ind)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "80")
    spacing.set(qn("w:line"), "300")
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.append(spacing)
    level.append(p_pr)
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE if marker_type != "check" else DARK_BLUE)
    r_pr.append(color)
    level.append(r_pr)
    abstract.append(level)
    numbering.append(abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def apply_num(paragraph, num_id):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num)
    p_pr.append(num_pr)


def add_inline(paragraph, text, *, base_size=11):
    parts = re.split(r"(`[^`]+`|\*\*[^*]+\*\*|\*[^*]+\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, "Consolas", 9.0, DARK_BLUE)
        elif part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, size=base_size, bold=True)
        elif part.startswith("*") and part.endswith("*"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, size=base_size, italic=True)
        else:
            run = paragraph.add_run(part)
            set_run_font(run, size=base_size, color=BLACK)


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ):
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def configure_page(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run("CbCR INDIA APPROACH  |  END-TO-END OPERATOR RUNBOOK")
    set_run_font(run, size=8.5, color=MUTED, bold=True)
    footer = section.footer
    p = footer.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(6.25))
    run = p.add_run("Implementation-aligned baseline  |  14 August 2026")
    set_run_font(run, size=8, color=MUTED)
    p.add_run("\tPage ")
    add_field(p, " PAGE ", "1")
    for run in p.runs:
        set_run_font(run, size=8, color=MUTED)


def add_title_block(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(42)
    p.paragraph_format.space_after = Pt(14)
    run = p.add_run("OPERATIONS GUIDE")
    set_run_font(run, size=10, color=BLUE, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run("CbCR Pipeline Operator Runbook")
    set_run_font(run, size=28, color=NAVY, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(20)
    run = p.add_run("A complete API and UI journey from source intake to DQ publication, risk results and case-selection decision")
    set_run_font(run, size=13, color=MUTED, italic=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(22)
    run = p.add_run("Version 1.0  |  FastAPI PoC 0.1.0  |  14 August 2026")
    set_run_font(run, size=9.5, color=DARK_BLUE, bold=True)
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.35)
    p.paragraph_format.right_indent = Inches(0.35)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(18)
    shade_paragraph(p, LIGHT_BLUE, BORDER)
    label = p.add_run("How to use this runbook. ")
    set_run_font(label, size=10.5, color=NAVY, bold=True)
    text = p.add_run("Execute the numbered steps in order for the ideal separated route. Capture every returned identifier, inspect the stated evidence, and apply the decision gate before moving to the next publication.")
    set_run_font(text, size=10.5, color=BLACK)


def table_widths(headers):
    count = len(headers)
    lower = [x.lower() for x in headers]
    if count == 2:
        return [2600, 6760]
    if count == 3 and lower[0] == "function":
        return [2500, 2800, 4060]
    if count == 3:
        return [1650, 2900, 4810]
    if count == 4 and lower[0] == "stage":
        return [750, 3150, 2300, 3160]
    if count == 4 and lower[0] == "profile":
        return [1800, 2300, 2360, 2900]
    if count == 4 and lower[0] in {"variable", "check", "symptom"}:
        return [1450, 2350, 2560, 3000]
    if count == 4:
        return [1400, 2300, 2700, 2960]
    return [USABLE_DXA // count] * (count - 1) + [USABLE_DXA - (USABLE_DXA // count) * (count - 1)]


def add_markdown_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    add_table_borders(table)
    hdr = table.rows[0]
    repeat_header(hdr)
    prevent_row_split(hdr)
    for idx, value in enumerate(headers):
        cell = hdr.cells[idx]
        set_cell_shading(cell, BLUE)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.08
        run = p.add_run(value)
        set_run_font(run, size=8.5, color=WHITE, bold=True)
    for row_index, values in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        for idx, value in enumerate(values):
            cell = row.cells[idx]
            if row_index % 2:
                set_cell_shading(cell, LIGHT_GRAY)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.08
            add_inline(p, value.replace("<br>", " / "), base_size=8.3)
    set_table_geometry(table, table_widths(headers))
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)


def add_code_block(doc, lines):
    for index, line in enumerate(lines or [""]):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.18)
        p.paragraph_format.right_indent = Inches(0.18)
        p.paragraph_format.space_before = Pt(4 if index == 0 else 0)
        p.paragraph_format.space_after = Pt(4 if index == len(lines) - 1 else 0)
        p.paragraph_format.line_spacing = 1.0
        p_pr = p._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), LIGHT_GRAY)
        p_pr.append(shd)
        if index == 0 or index == len(lines) - 1:
            p_bdr = OxmlElement("w:pBdr")
            edge = "top" if index == 0 else "bottom"
            line_border = OxmlElement(f"w:{edge}")
            line_border.set(qn("w:val"), "single")
            line_border.set(qn("w:sz"), "4")
            line_border.set(qn("w:color"), BORDER)
            p_bdr.append(line_border)
            p_pr.append(p_bdr)
        run = p.add_run(line or " ")
        set_run_font(run, "Consolas", 8.1, BLACK)


def build_document():
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    doc = Document()
    configure_styles(doc)
    configure_page(doc)
    add_title_block(doc)
    in_code = False
    code_lines = []
    active_list_type = None
    active_num_id = None
    index = 0
    while index < len(lines):
        line = lines[index]
        if line.startswith("# "):
            index += 1
            continue
        if line.startswith("```"):
            if not in_code:
                in_code = True
                code_lines = []
            else:
                add_code_block(doc, code_lines)
                in_code = False
            active_list_type = None
            active_num_id = None
            index += 1
            continue
        if in_code:
            code_lines.append(line)
            index += 1
            continue
        if line.startswith("| ") and index + 1 < len(lines) and re.match(r"^\|[-: |]+\|$", lines[index + 1]):
            headers = [x.strip() for x in line.strip("|").split("|")]
            rows = []
            index += 2
            while index < len(lines) and lines[index].startswith("|"):
                rows.append([x.strip() for x in lines[index].strip("|").split("|")])
                index += 1
            add_markdown_table(doc, headers, rows)
            active_list_type = None
            active_num_id = None
            continue
        if line.startswith("## "):
            p = doc.add_heading(line[3:], level=1)
            set_keep_with_next(p)
            active_list_type = None
            active_num_id = None
        elif line.startswith("### "):
            p = doc.add_heading(line[4:], level=2)
            set_keep_with_next(p)
            active_list_type = None
            active_num_id = None
        elif line.startswith("#### "):
            p = doc.add_heading(line[5:], level=3)
            set_keep_with_next(p)
            active_list_type = None
            active_num_id = None
        elif line.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.18)
            p.paragraph_format.right_indent = Inches(0.18)
            p.paragraph_format.space_before = Pt(5)
            p.paragraph_format.space_after = Pt(8)
            shade_paragraph(p)
            add_inline(p, line[2:], base_size=10.2)
            active_list_type = None
            active_num_id = None
        elif re.match(r"^- \[ \] ", line):
            if active_list_type != "check":
                active_list_type = "check"
                active_num_id = add_numbering(doc, "check")
            p = doc.add_paragraph()
            apply_num(p, active_num_id)
            add_inline(p, line[6:])
        elif line.startswith("- "):
            if active_list_type != "bullet":
                active_list_type = "bullet"
                active_num_id = add_numbering(doc, "bullet")
            p = doc.add_paragraph()
            apply_num(p, active_num_id)
            add_inline(p, line[2:])
        elif re.match(r"^\d+\. ", line):
            if active_list_type != "decimal":
                active_list_type = "decimal"
                active_num_id = add_numbering(doc, "decimal")
            p = doc.add_paragraph()
            apply_num(p, active_num_id)
            add_inline(p, re.sub(r"^\d+\. ", "", line))
        elif line.strip():
            p = doc.add_paragraph()
            add_inline(p, line)
            active_list_type = None
            active_num_id = None
        else:
            active_list_type = None
            active_num_id = None
        index += 1
    doc.core_properties.title = "CbCR Pipeline Operator Runbook"
    doc.core_properties.subject = "End-to-end database-first CbCR API and UI user journey"
    doc.core_properties.author = "CbCR India Approach"
    doc.core_properties.keywords = "CbCR, data quality, enrichment, risk, selection, API, runbook"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
