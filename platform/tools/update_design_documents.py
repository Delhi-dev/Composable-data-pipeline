from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"C:\Users\177885\Desktop\UP_West\TIWB\Work\CbCR_India_Approach")
DOCS = ROOT / "Design_Documents_DB_First"
PLATFORM_DOCS = ROOT / "cbcr-pipeline-scaffold" / "cbcr-pipeline-scaffold" / "platform" / "docs"
BLUE = "2E74B5"
NAVY = "1F4D78"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F4F6F9"
WHITE = "FFFFFF"
MUTED = "667085"
BLACK = "000000"
USABLE_DXA = 9360


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_width(table, widths: list[int]) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = Inches(widths[index] / 1440)
            tc_w = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
            tc_w.set(qn("w:w"), str(widths[index]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def style_run(run, *, name="Calibri", size=None, bold=None, italic=None, color=None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def setup_document(doc: Document, title: str, subtitle: str) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.18
    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 11.5, NAVY, 8, 4),
    ):
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.text = "CbCR India Approach  |  Database-first framework"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    style_run(header.runs[0], size=8.5, color=MUTED)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run("Implementation-aligned baseline  |  13 August 2026  |  Page ")
    style_run(r, size=8.5, color=MUTED)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    footer._p.append(fld)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(26)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run("CBCR INDIA APPROACH")
    style_run(r, size=10, bold=True, color=BLUE)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    r = p.add_run(title)
    style_run(r, size=25, bold=True, color=BLACK)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run(subtitle)
    style_run(r, size=12.5, color=MUTED)
    add_callout(doc, "Document status", "Authoritative implementation-aligned baseline. The formal FDD, TDD and sequence documents should be read together with this guide.")


def add_callout(doc: Document, label: str, text: str, fill=LIGHT_BLUE) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    set_table_width(table, [USABLE_DXA])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(f"{label}. ")
    style_run(r, bold=True, color=NAVY)
    r = p.add_run(text)
    style_run(r, color=BLACK)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[int] | None = None) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    if widths is None:
        base = USABLE_DXA // len(headers)
        widths = [base] * (len(headers) - 1) + [USABLE_DXA - base * (len(headers) - 1)]
    set_table_width(table, widths)
    set_repeat_table_header(table.rows[0])
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = header
        set_cell_shading(cell, BLUE)
        for run in cell.paragraphs[0].runs:
            style_run(run, size=9, bold=True, color=WHITE)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(values):
            cells[i].text = str(value)
            if row_index % 2:
                set_cell_shading(cells[i], LIGHT_GRAY)
            for paragraph in cells[i].paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    style_run(run, size=8.6)
    set_table_width(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_bullets(doc: Document, values: list[str], numbered=False) -> None:
    style = "List Number" if numbered else "List Bullet"
    for value in values:
        p = doc.add_paragraph(style=style)
        p.paragraph_format.space_after = Pt(3)
        p.add_run(value)


def replace_all(doc: Document, old: str, new: str) -> None:
    for paragraph in doc.paragraphs:
        if old in paragraph.text:
            for run in paragraph.runs:
                if old in run.text:
                    run.text = run.text.replace(old, new)
                    break
            else:
                paragraph.text = paragraph.text.replace(old, new)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    if old in paragraph.text:
                        for run in paragraph.runs:
                            if old in run.text:
                                run.text = run.text.replace(old, new)
                                break


def append_alignment_header(doc: Document, heading: str) -> None:
    doc.add_page_break()
    doc.add_heading(heading, level=1)
    add_callout(
        doc,
        "Status convention",
        "Implemented and exercised means available in the running PoC. Applied database contract means present in PostgreSQL DDL but not necessarily used yet by the FastAPI persistence layer. Scaffolded and planned items still require production implementation.",
    )


def update_fdd() -> None:
    path = DOCS / "CbCR_DB_First_Functional_Design_Document.docx"
    doc = Document(path)
    replace_all(doc, "Database ingestion first; XML is explicitly deferred", "Database ingestion first; XML lossless schema prepared and XML parser deferred")
    append_alignment_header(doc, "9. Implementation alignment update - 13 August 2026")
    doc.add_paragraph(
        "The functional baseline is now demonstrated against the local PostgreSQL CbCR store. Four live reports, 7,260 constituent entities and 309 jurisdiction summaries traverse the physical-to-canonical boundary. Existing business requirements remain valid; the table below records delivery status and the remaining production gap."
    )
    doc.add_heading("9.1 Implemented functional lifecycle", level=2)
    add_bullets(doc, [
        "Register fixture or PostgreSQL sources through editable endpoint links and secret references; credentials are not persisted.",
        "Map different physical names into one report/entity/jurisdiction-metric vocabulary and expose the field dictionary to authorised users.",
        "Create immutable landing and canonical datasets, select independent functions, supply UI-generated parameters and receive asynchronous run status.",
        "Persist findings, evaluate blocking severities, disposition every report and publish only eligible records into a new DQ-passed dataset.",
        "Publish enrichment attributes, create the risk-ready dataset, persist risk results/score components, generate candidates and record human decisions.",
        "Derive access from multiple roles and audit source, mapping, run, promotion, catalogue and decision actions.",
    ])
    doc.add_heading("9.2 Functional status matrix", level=2)
    add_table(doc, ["Capability", "Status", "Functional note"], [
        ["PostgreSQL CbCR input", "Implemented", "Live source view and mapping exercised end to end."],
        ["Single/filtered/bulk scope", "Implemented", "Optional report identifiers; absence means all source records."],
        ["Independent selectable checks", "Implemented", "Fourteen sample functions; one child execution per selection."],
        ["Asynchronous checks", "Implemented PoC", "HTTP 202 and status polling; worker is currently in-process."],
        ["Asynchronous intake", "Planned", "Lot intake is currently synchronous and must become queued."],
        ["DQ gate and filtered publication", "Implemented", "Per-report eligible/rejected dispositions and copied descendants."],
        ["Enrichment/risk-ready/risk/selection", "Implemented PoC", "Versioned datasets and persisted evidence/actions."],
        ["Many-role user access", "Implemented PoC", "Central action/stage/purpose checks; header identity is non-production."],
        ["XML intake", "Schema prepared", "Lossless PostgreSQL structures exist; parser and full XSD bundle remain."],
        ["PostgreSQL control persistence", "Planned next", "Running control plane remains SQLite; PostgreSQL contract is applied."],
    ], [2200, 1600, 5560])
    doc.add_heading("9.3 Business-data limitation discovered", level=2)
    doc.add_paragraph(
        "The former XML parser did not retain currency attributes for the 309 existing jurisdiction summaries. The framework exposes these values as missing so data-quality functions can identify the defect; it does not invent values. Recovery requires reprocessing the original XML. Future loads can populate a currency value for every monetary measure."
    )
    doc.add_heading("9.4 Revised acceptance evidence", level=2)
    add_bullets(doc, [
        "All four local PostgreSQL reports map into the stable canonical contract.",
        "Live intake creates four reports, 7,260 entities and 309 metrics in governed PoC datasets.",
        "Three idempotent PostgreSQL migrations and application-role access were verified.",
        "Twenty-one automated tests cover RBAC, portability, asynchronous independence, lifecycle, immutability and adapter safety.",
    ])
    doc.add_heading("9.5 Outstanding business decisions", level=2)
    add_bullets(doc, [
        "Approval of DQ blocking policy, enrichment sources, risk formulas and candidate-decision workflow.",
        "Correction/deletion treatment, retention periods, report templates and permitted bulk exports.",
        "Production identity provider, role ownership, segregation of duties and user recertification.",
        "Whether trusted upstream DQ/enrichment profiles may omit separate physical publications.",
    ])
    doc.save(path)


def update_tdd() -> None:
    path = DOCS / "CbCR_DB_First_Technical_Design_Document.docx"
    doc = Document(path)
    replace_all(doc, "Database ingestion first; XML is explicitly deferred", "Database ingestion first; XML lossless schema prepared and XML parser deferred")
    append_alignment_header(doc, "15. As-built architecture and database design")
    doc.add_heading("15.1 Current runtime versus prepared production contract", level=2)
    add_table(doc, ["Concern", "Executable PoC", "Applied PostgreSQL contract / target"], [
        ["Source data", "Fixture JSON or local PostgreSQL canonical view", "Populated public.cbcr_* XML model retained and extended"],
        ["Control persistence", "SQLite", "cbcr_control tables with mapping, queue, RBAC, datasets, lineage and audit"],
        ["Business zones", "SQLite stg/can/dqp/enr/rr table families", "Separated cbcr_canonical, cbcr_dq_published, cbcr_enriched and cbcr_risk_ready schemas"],
        ["Worker", "In-process asyncio poller", "Durable broker-backed stateless workers using function_run leases"],
        ["Identity", "x-cbcr-user demonstration header", "Jurisdiction IdP/OIDC plus mapped RBAC claims"],
    ], [1800, 3200, 4360])
    doc.add_heading("15.2 PostgreSQL schema map", level=2)
    add_table(doc, ["Schema", "Responsibility", "Representative objects"], [
        ["public", "Existing XML-consumption source", "message, body, reporting entity, reports, summaries, entities, repeatable XSD children"],
        ["cbcr_staging", "Exact XML and validation evidence", "xml_document, xml_validation_issue"],
        ["cbcr_control", "Metadata, queue, RBAC and lineage", "source, mapping, function, lot, dataset, run/function_run, user/role, audit"],
        ["cbcr_canonical", "Stable mapped contract", "report/entity/metric plus v_xml_* and v_api_report_json"],
        ["cbcr_dq", "Quality evidence and gate", "finding, gate evaluation, disposition, promotion event"],
        ["cbcr_dq_published", "Filtered publication", "report, entity, jurisdiction_metric"],
        ["cbcr_enriched", "Enriched publication", "common grains plus attribute"],
        ["cbcr_risk_ready", "Final risk input", "report, entity, jurisdiction_metric"],
        ["cbcr_risk", "Risk output", "result, score_component"],
        ["cbcr_selection", "Candidates and decisions", "selection_batch, candidate, decision_history"],
    ], [1550, 3200, 4610])
    doc.add_heading("15.3 XSD-completeness extensions", level=2)
    add_bullets(doc, [
        "Exact XML landing, schema-validation issues and parser status/version evidence.",
        "Repeatable CbcBody, receiving countries, message/document correction references and organisation residences/names/identifiers.",
        "Structured AddressFix fields, legal-address type and repeatable constituent business activities.",
        "Independent currency attributes for all nine monetary measures.",
        "Repeatable AdditionalInfo text, countries and summary references.",
    ])
    doc.add_heading("15.4 Canonical adapter and safety", level=2)
    doc.add_paragraph(
        "The postgresql-canonical adapter reads only report_payload from the allow-listed cbcr_canonical.v_api_report_json view. Its URL must contain user, host, database and an identifier-safe schema-qualified view. Embedded passwords and arbitrary options are rejected. The secret is resolved only from an env:NAME reference. Rules receive CanonicalDataAPI, not a connection or SQL cursor."
    )
    doc.add_heading("15.5 Applied integrity", level=2)
    add_bullets(doc, [
        "Foreign keys bind mapping versions, dataset parents/children, evaluations, promotions, enrichment, risk and selections.",
        "Report/entity/metric rows bind to a registered dataset version and child rows bind to their report.",
        "Twenty-four triggers reject update/delete across all four published business-data schemas.",
        "Raw XML payload and hash are immutable while parse-status metadata may progress.",
    ])
    doc.add_heading("15.6 Active HTTP surface", level=2)
    add_table(doc, ["API family", "Principal operations"], [
        ["Sources and mappings", "List/save source; draft/validate/activate mapping; canonical and physical dictionaries"],
        ["Pipeline and data lifecycle", "Profiles; lots; datasets; reports; lineage; evaluate; dispositions; promote"],
        ["Function catalogue", "List; create draft; activate; retire; delete draft"],
        ["Runs and evidence", "Queue source/dataset run; status; findings; risk results"],
        ["Selection", "Batches; candidates; human decision and rationale"],
        ["Access and audit", "Users; roles; replace many-role assignment; audit events"],
    ], [2300, 7060])
    doc.add_heading("15.7 Verified technical evidence and limitations", level=2)
    add_bullets(doc, [
        "Three PostgreSQL migrations applied and re-run safely; no source table or row was deleted.",
        "The cbcr_user application identity reads the canonical views and 25-entry dictionary.",
        "Live canonical mapping and end-to-end intake counts reconcile to 4 reports, 7,260 entities and 309 metrics.",
        "Twenty-one tests pass. Production work remains for PostgreSQL repositories, queued ingestion, durable workers, IdP, secret manager, operational controls and the XML parser.",
    ])
    doc.save(path)


def update_sequence() -> None:
    path = DOCS / "CbCR_DB_First_Implementation_Sequence.docx"
    doc = Document(path)
    replace_all(doc, "Database ingestion first; XML is explicitly deferred", "Database ingestion first; XML lossless schema prepared and XML parser deferred")
    append_alignment_header(doc, "10. Delivery progress and revised sequence")
    doc.add_heading("10.1 Completed implementation slice", level=2)
    add_table(doc, ["Work package", "Status", "Evidence"], [
        ["Canonical models and portability", "Complete PoC", "Two fixture schemas plus PostgreSQL map through one rule contract"],
        ["Source/mapping/dictionary catalogue", "Complete PoC", "Editable sources; versioned mappings; 25 PostgreSQL dictionary rows"],
        ["Data lifecycle", "Complete PoC", "Landing through risk-ready, gated promotion, enrichment and lineage"],
        ["Independent asynchronous functions", "Complete PoC", "One queued child per function; isolated failure/result"],
        ["UI and RBAC", "Complete PoC", "Run/lifecycle/catalogue/access workspaces; many-role union and deny precedence"],
        ["Risk and selection persistence", "Complete PoC", "Results/components, candidates and human decisions"],
        ["PostgreSQL source integration", "Complete", "Live adapter and 4/7,260/309 count reconciliation"],
        ["PostgreSQL framework DDL", "Applied contract", "Separated schemas, queue/RBAC, FKs and 24 immutability triggers"],
        ["XML lossless model", "Applied contract", "Root v2.0 cardinality structures prepared; parser pending"],
    ], [2400, 1700, 5260])
    doc.add_heading("10.2 Next critical chain", level=2)
    add_bullets(doc, [
        "Implement PostgreSQL repositories for the active control/data plane and migrate the SQLite PoC state model.",
        "Queue lot/source intake and add paging, source watermark, snapshot isolation, reconciliation and restartability.",
        "Deploy durable broker-backed workers with leases, bounded retry, idempotency and operational recovery.",
        "Implement the XML v2.0 loader using the complete official imported XSD bundle and populate all repeatable child structures.",
        "Split database identities/grants for parser, control service, DQ, risk and read-only UI; connect the jurisdiction secret manager.",
        "Replace demonstration identity with IdP/OIDC and implement role lifecycle/recertification.",
        "Add monitoring, backup/restore, retention, privacy, performance and failover testing.",
        "Onboard and business-approve production DQ/risk rules through the existing function lifecycle.",
    ], numbered=True)
    doc.add_heading("10.3 Exit gates for the next release", level=2)
    add_table(doc, ["Gate", "Required evidence"], [
        ["PostgreSQL parity", "All current 21 tests plus lifecycle integration run against PostgreSQL repositories"],
        ["Durability", "API and workers restart without lost/duplicated effect; leases and retry demonstrated"],
        ["XML fidelity", "Official valid/invalid/correction fixtures reconcile to XML and canonical counts"],
        ["Security", "IdP, secret injection, least privilege, audit export and access recertification approved"],
        ["Operations", "Dashboards, alerts, backup restore, retention and recovery objectives exercised"],
        ["Business", "DQ gate, enrichment provenance, risk rules and selection workflow signed off"],
    ], [2100, 7260])
    doc.add_heading("10.4 Dependencies requiring jurisdiction input", level=2)
    add_bullets(doc, [
        "Complete official XML XSD dependency bundle and representative original XML files for backfill/validation.",
        "Identity-provider and secret-manager integration pattern.",
        "Production database/network topology, environments, retention and backup standards.",
        "Approved business rule definitions, thresholds, purposes, roles and promotion policy.",
    ])
    doc.save(path)


def add_markdown_inline(paragraph, text: str) -> None:
    parts = re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text)
    for part in parts:
        if part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            style_run(run, name="Consolas", size=8.8, color=NAVY)
        elif part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        else:
            paragraph.add_run(part)


def markdown_to_docx(markdown_path: Path, output_path: Path) -> None:
    lines = markdown_path.read_text(encoding="utf-8").splitlines()
    doc = Document()
    setup_document(
        doc,
        "Schemas, Layers and APIs Guide",
        "Complete operating and technical reference for the CbCR DQ, enrichment, risk and selection framework",
    )
    in_code = False
    code_lines: list[str] = []
    index = 0
    while index < len(lines):
        line = lines[index]
        if line.startswith("# "):
            index += 1
            continue
        if line.startswith("```"):
            if not in_code:
                in_code = True
                code_lines = []
            else:
                table = doc.add_table(rows=1, cols=1)
                table.style = "Table Grid"
                set_table_width(table, [USABLE_DXA])
                cell = table.cell(0, 0)
                set_cell_shading(cell, LIGHT_GRAY)
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                r = p.add_run("\n".join(code_lines))
                style_run(r, name="Consolas", size=8.1)
                in_code = False
            index += 1
            continue
        if in_code:
            code_lines.append(line)
            index += 1
            continue
        if line.startswith("## "):
            doc.add_heading(line[3:], level=1)
        elif line.startswith("### "):
            doc.add_heading(line[4:], level=2)
        elif line.startswith("#### "):
            doc.add_heading(line[5:], level=3)
        elif line.startswith("| ") and index + 1 < len(lines) and re.match(r"^\|[-: |]+\|$", lines[index + 1]):
            headers = [value.strip() for value in line.strip("|").split("|")]
            rows = []
            index += 2
            while index < len(lines) and lines[index].startswith("|"):
                rows.append([value.strip().replace("<br>", " / ") for value in lines[index].strip("|").split("|")])
                index += 1
            if len(headers) == 2:
                widths = [2300, 7060]
            elif len(headers) == 3:
                widths = [2100, 2200, 5060]
            elif len(headers) == 4:
                widths = [1250, 2200, 2950, 2960]
            else:
                widths = None
            add_table(doc, headers, rows, widths)
            continue
        elif re.match(r"^- ", line):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(3)
            add_markdown_inline(p, line[2:])
        elif re.match(r"^\d+\. ", line):
            # Keep the Markdown number literal so each independently authored
            # sequence restarts exactly where the source document says it does.
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.28)
            p.paragraph_format.first_line_indent = Inches(-0.28)
            p.paragraph_format.space_after = Pt(3)
            number, value = line.split(". ", 1)
            p.add_run(f"{number}.\t")
            add_markdown_inline(p, value)
        elif line.strip():
            p = doc.add_paragraph()
            add_markdown_inline(p, line)
        index += 1
    doc.core_properties.title = "CbCR Framework - Schemas, Layers and APIs Guide"
    doc.core_properties.subject = "Implementation-aligned database and API architecture"
    doc.core_properties.author = "CbCR India Approach"
    doc.save(output_path)


def main() -> None:
    DOCS.mkdir(parents=True, exist_ok=True)
    update_fdd()
    update_tdd()
    update_sequence()
    markdown_to_docx(
        PLATFORM_DOCS / "COMPLETE_FRAMEWORK_GUIDE.md",
        DOCS / "CbCR_Framework_Schemas_Layers_APIs_Guide.docx",
    )


if __name__ == "__main__":
    main()
